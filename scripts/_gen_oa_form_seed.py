# -*- coding: utf-8 -*-
"""
_gen_oa_form_seed.py — docs/52 · P5 首批 43 个 OA 模块表单定义种子生成器
==================================================================================

用途
----
读取《OA系统功能清单_功能描述补充版.xlsx》拿到 43 个（模块名 / 描述 / 来源），
把每个模块映射到 agent_worker.worker_type 的稳定码（module_code），为每个模块设计
3–6 个字段的 JSON Schema（Draft 2020-12），生成**幂等、可复跑**的迁移 SQL：

    server/aioa-boot/src/main/resources/db/migration/V86__oa_module_forms.sql

设计口径（与 docs/52 §3.2 / §3.3 / V85 契约一致）
-------------------------------------------------
1. 字段类型白名单（冻结，禁止新造词）：
   text | textarea | number | amount | date | datetime | select | multiselect | file | boolean
2. schema_json 同时供「LLM 参数定义 / H5 渲染 / 后端校验」三处同源。
   每个属性用 JSON Schema 原生 type（string/number/boolean/array）承载 JSON 类型，
   用扩展关键字 `x-type` 承载上面的白名单控件类型（渲染器据此选控件）；
   select/multiselect 用 `enum` + `enumNames` 内联静态选项（docs/52 §3.2 既定）。
3. biz_type：**只有请假用 'LEAVE'**，其余一律为空（其余模块没有 ACTIVE 审批流，
   乱填会让提交直接失败 —— 这是 lead 的硬约束）。
4. submit_policy：请假 = 'APPROVAL+RESUME'；其余 = 'LEDGER'。
5. worker_capability：每个模块 1 条能力声明；从 agent_worker 派生（模板 tenant=0
   + 各租户实例），用 INSERT...SELECT + WHERE NOT EXISTS 保证幂等。
6. form_template_version：每个 form_template 行（平台 tenant=0 + 各租户回填）各一条
   version=1 快照，幂等。

字段推导诚实声明
----------------
- `derived=True`：字段名/类型直接来自 xlsx 描述里的动作动词（申请/登记/填报/上报/
  借阅/申诉/起草/预约/发布/立项/录入/建档…），属于严格推导。
- `derived=False`：平台/能力型模块（门户首页、审批中心、系统集成、数据分析、移动办公），
  描述讲的是平台能力而非单条可填业务，字段为合理默认。
- 无论 derived 真假，**select/multiselect 的枚举值都是合理默认**：描述只点名类别
  （如「按紧急度分级」），未穷举具体选项，枚举由本脚本给出。

复跑
----
    python scripts/_gen_oa_form_seed.py
生成结果写入 V86 文件（路径见下方 OUT_SQL）。不改 xlsx、不改任何已有迁移。

依赖：openpyxl（仅读取 xlsx 校验行数，字段设计已固化为下方 MODULES，不依赖 xlsx 内容）。
"""
import json
import os
import sys

# ---------------------------------------------------------------------------
# 路径
# ---------------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
XLSX = r"C:\Users\刘尖尖\Desktop\OA系统功能清单_功能描述补充版.xlsx"
OUT_SQL = os.path.join(
    REPO,
    "server", "aioa-boot", "src", "main", "resources", "db", "migration",
    "V86__oa_module_forms.sql",
)

# 字段类型白名单（docs/52 §3.2 冻结）
TYPE_WHITELIST = {
    "text", "textarea", "number", "amount", "date", "datetime",
    "select", "multiselect", "file", "boolean",
}

# JSON Schema 原生 type（与 x-type 白名单控件对应）
BASE_JSON_TYPE = {
    "text": "string", "textarea": "string", "date": "string", "datetime": "string",
    "select": "string", "file": "string",
    "number": "number", "amount": "number",
    "boolean": "boolean",
    "multiselect": "array",
}


def opt(pairs):
    """构造 select/multiselect 的枚举选项：[(value, label), ...]"""
    return [{"value": v, "label": l} for v, l in pairs]


# ---------------------------------------------------------------------------
# 43 个模块的能力声明 + 字段设计（seq 与 xlsx 序号 1..43 严格对齐）
# 每个模块：name / source / worker_type / module_code / form_code /
#           task_name / biz_type / submit_policy / derived / fields
# fields: [{name, label, type, required, placeholder?, options?}]
# ---------------------------------------------------------------------------
MODULES = [
    # ---- 原有 26 ----
    dict(seq=1, name="公文助手", source="原有", worker_type="OA_DOC", module_code="doc",
         form_code="doc_draft", task_name="公文起草", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="doc_kind", label="公文种类", type="select", required=True,
                  options=opt([("OUT", "发文"), ("IN", "收文")])),
             dict(name="title", label="公文标题", type="text", required=True, placeholder="请输入公文标题"),
             dict(name="urgency", label="紧急程度", type="select", required=True,
                  options=opt([("NORMAL", "普通"), ("URGENT", "紧急"), ("FLASH", "特急")])),
             dict(name="secret_level", label="涉密等级", type="select", required=False,
                  options=opt([("PUBLIC", "非涉密"), ("INTERNAL", "内部"), ("SECRET", "秘密"), ("CONFIDENTIAL", "机密")])),
             dict(name="summary", label="正文摘要", type="textarea", required=True, placeholder="请填写公文正文要点"),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=2, name="档案助手", source="原有", worker_type="OA_ARCHIVE", module_code="archive",
         form_code="archive_borrow", task_name="档案借阅", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="archive_type", label="档案类别", type="select", required=True,
                  options=opt([("DOC", "文书档案"), ("TECH", "科技档案"), ("ACCT", "会计档案"), ("AV", "声像档案")])),
             dict(name="archive_no", label="档号", type="text", required=False, placeholder="已知档号可填，便于调阅"),
             dict(name="title", label="档案题名", type="text", required=True, placeholder="请输入档案题名"),
             dict(name="borrow_use", label="借阅用途", type="textarea", required=True, placeholder="说明借阅用途"),
             dict(name="return_date", label="预计归还日期", type="date", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=3, name="会议助手", source="原有", worker_type="OA_MEETING", module_code="meeting",
         form_code="meeting_book", task_name="会议预约", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="title", label="会议主题", type="text", required=True, placeholder="请输入会议主题"),
             dict(name="meeting_type", label="会议类型", type="select", required=True,
                  options=opt([("REGULAR", "例会"), ("THEME", "专题会"), ("DECISION", "决策会"), ("TRAIN", "培训会")])),
             dict(name="start_time", label="开始时间", type="datetime", required=True),
             dict(name="end_time", label="结束时间", type="datetime", required=True),
             dict(name="attendees", label="参会人", type="textarea", required=True, placeholder="列明参会人员/部门"),
             dict(name="agenda", label="议程", type="textarea", required=False),
         ]),
    dict(seq=4, name="通讯录", source="原有", worker_type="OA_CONTACT", module_code="contact",
         form_code="contact_add", task_name="联系人维护", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="person_name", label="姓名", type="text", required=True),
             dict(name="dept", label="部门", type="text", required=True),
             dict(name="title", label="职务", type="text", required=False),
             dict(name="phone", label="手机号", type="text", required=True, placeholder="联系电话"),
             dict(name="email", label="邮箱", type="text", required=False),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=5, name="工作日志", source="原有", worker_type="OA_WORKLOG", module_code="worklog",
         form_code="worklog_submit", task_name="工作日志填报", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="log_date", label="日期", type="date", required=True),
             dict(name="today_work", label="今日工作", type="textarea", required=True, placeholder="记录当日完成工作"),
             dict(name="tomorrow_plan", label="明日计划", type="textarea", required=False),
             dict(name="project_tag", label="关联项目", type="text", required=False, placeholder="可选关联项目"),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=6, name="工作计划", source="原有", worker_type="OA_WORKPLAN", module_code="workplan",
         form_code="workplan_submit", task_name="工作计划填报", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="plan_name", label="计划名称", type="text", required=True),
             dict(name="cycle", label="计划周期", type="select", required=True,
                  options=opt([("YEAR", "年度"), ("MONTH", "月度"), ("WEEK", "周")])),
             dict(name="goal", label="目标", type="textarea", required=True),
             dict(name="owner", label="责任人", type="text", required=True),
             dict(name="start_date", label="开始日期", type="date", required=True),
             dict(name="end_date", label="结束日期", type="date", required=True),
         ]),
    dict(seq=7, name="工作周报", source="原有", worker_type="OA_WEEKLY", module_code="weekly",
         form_code="weekly_submit", task_name="工作周报填报", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="week_no", label="周次", type="text", required=True, placeholder="如 2026-W42"),
             dict(name="done", label="本周完成", type="textarea", required=True),
             dict(name="next_plan", label="下周计划", type="textarea", required=True),
             dict(name="blockers", label="需协调事项", type="textarea", required=False),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=8, name="费用报销", source="原有", worker_type="OA_EXPENSE", module_code="expense",
         form_code="expense_form", task_name="费用报销申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="expense_type", label="费用类型", type="select", required=True,
                  options=opt([("TRAVEL", "差旅"), ("ENTERTAIN", "招待"), ("OFFICE", "办公"), ("OTHER", "其他")])),
             dict(name="amount", label="金额", type="amount", required=True),
             dict(name="occur_date", label="发生日期", type="date", required=True),
             dict(name="reason", label="费用说明", type="textarea", required=True),
             dict(name="invoice", label="发票", type="file", required=True),
             dict(name="project", label="所属项目", type="text", required=False),
         ]),
    dict(seq=9, name="出差助手", source="原有", worker_type="OA_TRAVEL", module_code="travel",
         form_code="travel_form", task_name="出差申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="destination", label="出差地点", type="text", required=True),
             dict(name="start_date", label="出发日期", type="date", required=True),
             dict(name="end_date", label="返回日期", type="date", required=True),
             dict(name="reason", label="出差事由", type="textarea", required=True),
             dict(name="transport", label="交通工具", type="select", required=True,
                  options=opt([("FLIGHT", "飞机"), ("RAIL", "高铁"), ("BUS", "汽车"), ("CAR", "自驾")])),
             dict(name="est_amount", label="预计费用", type="amount", required=False),
         ]),
    dict(seq=10, name="印章助手", source="原有", worker_type="OA_SEAL", module_code="seal",
         form_code="seal_apply", task_name="用印申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="seal_type", label="印章类型", type="select", required=True,
                  options=opt([("OFFICIAL", "公章"), ("CONTRACT", "合同章"), ("FINANCE", "财务章"), ("LEGAL", "法人章")])),
             dict(name="file_title", label="用印文件标题", type="text", required=True),
             dict(name="reason", label="用印事由", type="textarea", required=True),
             dict(name="use_date", label="用印日期", type="date", required=True),
             dict(name="copies", label="用印份数", type="number", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=11, name="项目助手", source="原有", worker_type="OA_PROJECT", module_code="project",
         form_code="project_create", task_name="项目立项", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="project_name", label="项目名称", type="text", required=True),
             dict(name="goal", label="项目目标", type="textarea", required=True),
             dict(name="start_date", label="开始日期", type="date", required=True),
             dict(name="end_date", label="结束日期", type="date", required=True),
             dict(name="budget", label="预算金额", type="amount", required=False),
             dict(name="risk", label="风险说明", type="textarea", required=False),
         ]),
    dict(seq=12, name="收藏夹", source="原有", worker_type="OA_FAVORITE", module_code="favorite",
         form_code="favorite_add", task_name="收藏添加", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="fav_type", label="收藏类型", type="select", required=True,
                  options=opt([("URL", "网址"), ("FILE", "文件"), ("IMAGE", "图片"), ("NOTE", "笔记")])),
             dict(name="title", label="标题", type="text", required=True),
             dict(name="content", label="链接或内容", type="textarea", required=True),
             dict(name="tags", label="标签", type="text", required=False, placeholder="逗号分隔"),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=13, name="大事记", source="原有", worker_type="OA_MILESTONE", module_code="milestone",
         form_code="milestone_add", task_name="大事记录入", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="event_title", label="事件标题", type="text", required=True),
             dict(name="event_date", label="事件日期", type="date", required=True),
             dict(name="category", label="事件分类", type="select", required=True,
                  options=opt([("HONOR", "荣誉"), ("MILESTONE", "里程碑"), ("ACTIVITY", "重大活动"), ("CHANGE", "变更")])),
             dict(name="description", label="事件描述", type="textarea", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=14, name="工作日历", source="原有", worker_type="OA_CALENDAR", module_code="calendar",
         form_code="calendar_event", task_name="日程新建", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="event_title", label="日程标题", type="text", required=True),
             dict(name="start_time", label="开始时间", type="datetime", required=True),
             dict(name="end_time", label="结束时间", type="datetime", required=True),
             dict(name="all_day", label="是否全天", type="boolean", required=False),
             dict(name="remind_way", label="提醒方式", type="select", required=False,
                  options=opt([("POP", "弹窗"), ("SMS", "短信"), ("MAIL", "邮件")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=15, name="通知公告", source="原有", worker_type="OA_NOTICE", module_code="notice",
         form_code="notice_publish", task_name="通知公告发布", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="title", label="标题", type="text", required=True),
             dict(name="category", label="公告分类", type="select", required=True,
                  options=opt([("NOTICE", "通知"), ("BULLETIN", "公告"), ("RULE", "制度"), ("HR", "人事")])),
             dict(name="scope", label="发布范围", type="textarea", required=True, placeholder="指定部门/人员或全员"),
             dict(name="body", label="正文", type="textarea", required=True),
             dict(name="pinned", label="是否置顶", type="boolean", required=False),
             dict(name="expire_date", label="有效期至", type="date", required=False),
         ]),
    dict(seq=16, name="信息发布", source="原有", worker_type="OA_PUBLISH", module_code="publish",
         form_code="publish_article", task_name="信息发布", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="column", label="栏目", type="select", required=True,
                  options=opt([("NEWS", "新闻动态"), ("SPECIAL", "专题"), ("HR", "人事风采"), ("CULTURE", "企业文化")])),
             dict(name="title", label="标题", type="text", required=True),
             dict(name="body", label="正文", type="textarea", required=True),
             dict(name="author", label="作者", type="text", required=False),
             dict(name="schedule_at", label="定时发布", type="datetime", required=False),
         ]),
    dict(seq=17, name="工作安排", source="原有", worker_type="OA_ASSIGN", module_code="assign",
         form_code="assign_task", task_name="工作安排下发", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="task_title", label="任务标题", type="text", required=True),
             dict(name="owner", label="责任人", type="text", required=True),
             dict(name="requirement", label="任务要求", type="textarea", required=True),
             dict(name="due_date", label="截止日期", type="date", required=True),
             dict(name="priority", label="优先级", type="select", required=True,
                  options=opt([("HIGH", "高"), ("MID", "中"), ("LOW", "低")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=18, name="重要提醒", source="原有", worker_type="OA_REMIND", module_code="remind",
         form_code="remind_rule", task_name="提醒规则设置", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="提醒事项", type="text", required=True),
             dict(name="remind_type", label="提醒类型", type="select", required=True,
                  options=opt([("APPROVAL", "审批"), ("TASK", "任务"), ("SCHEDULE", "日程"),
                               ("CONTRACT", "合同到期"), ("LICENSE", "证照到期"), ("DEVICE", "设备维保")])),
             dict(name="remind_time", label="提醒时间", type="datetime", required=True),
             dict(name="urgency", label="紧急程度", type="select", required=True,
                  options=opt([("URGENT", "紧急"), ("NORMAL", "普通")])),
             dict(name="channels", label="提醒渠道", type="multiselect", required=False,
                  options=opt([("POP", "弹窗"), ("SMS", "短信"), ("MAIL", "邮件")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=19, name="用车助手", source="原有", worker_type="OA_VEHICLE", module_code="vehicle",
         form_code="vehicle_apply", task_name="用车申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="reason", label="用车事由", type="textarea", required=True),
             dict(name="from_addr", label="出发地", type="text", required=True),
             dict(name="to_addr", label="目的地", type="text", required=True),
             dict(name="start_time", label="用车开始", type="datetime", required=True),
             dict(name="end_time", label="用车结束", type="datetime", required=True),
             dict(name="passengers", label="乘车人数", type="number", required=True),
         ]),
    dict(seq=20, name="请假助手", source="原有", worker_type="OA_LEAVE", module_code="leave",
         form_code="leave_request", task_name="请假申请", biz_type="LEAVE", submit_policy="APPROVAL+RESUME", derived=True,
         fields=[
             dict(name="leave_type", label="假种", type="select", required=True,
                  options=opt([("ANNUAL", "年休假"), ("SICK", "病假"), ("PERSONAL", "事假"),
                               ("COMPENSATORY", "调休"), ("MARRIAGE", "婚假"), ("MATERNITY", "产假")])),
             dict(name="start_date", label="开始日期", type="date", required=True),
             dict(name="end_date", label="结束日期", type="date", required=True),
             dict(name="days", label="请假天数", type="number", required=True),
             dict(name="reason", label="请假事由", type="textarea", required=True),
             dict(name="proof", label="证明材料", type="file", required=False),
         ]),
    dict(seq=21, name="考勤管理", source="原有", worker_type="OA_ATTENDANCE", module_code="attendance",
         form_code="attendance_appeal", task_name="考勤异常申诉", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="appeal_type", label="申诉类型", type="select", required=True,
                  options=opt([("MISS", "漏打卡"), ("LATE", "迟到"), ("EARLY", "早退"), ("OUT", "外勤")])),
             dict(name="appeal_date", label="申诉日期", type="date", required=True),
             dict(name="description", label="说明", type="textarea", required=True),
             dict(name="proof", label="证明材料", type="file", required=False),
         ]),
    dict(seq=22, name="文档管理", source="原有", worker_type="OA_DOCMGR", module_code="docmgr",
         form_code="doc_upload", task_name="文档上传", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="doc_title", label="文档标题", type="text", required=True),
             dict(name="doc_category", label="文档分类", type="select", required=True,
                  options=opt([("POLICY", "制度"), ("TECH", "技术"), ("REPORT", "报告"), ("TEMP", "模板")])),
             dict(name="share_scope", label="共享范围", type="textarea", required=True, placeholder="指定部门/人员"),
             dict(name="expire_date", label="有效期至", type="date", required=False),
             dict(name="file", label="文件", type="file", required=True),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=23, name="党务助手", source="原有", worker_type="OA_PARTY", module_code="party",
         form_code="party_apply", task_name="党务事项办理", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="事项类型", type="select", required=True,
                  options=opt([("FEE", "党费缴纳"), ("LIFESTUDY", "组织生活报名"), ("DEV", "发展流程申报")])),
             dict(name="person_name", label="姓名", type="text", required=True),
             dict(name="amount", label="金额", type="amount", required=False),
             dict(name="description", label="说明", type="textarea", required=False),
         ]),
    dict(seq=24, name="工会工作", source="原有", worker_type="OA_UNION", module_code="union",
         form_code="union_apply", task_name="工会事项办理", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="事项类型", type="select", required=True,
                  options=opt([("HELP", "困难帮扶"), ("ACTIVITY", "活动报名"), ("PROPOSAL", "提案征集"), ("MODEL", "劳模评选")])),
             dict(name="applicant", label="申请人", type="text", required=True),
             dict(name="description", label="说明", type="textarea", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=25, name="待办助手", source="原有", worker_type="OA_TODO", module_code="todo",
         form_code="todo_add", task_name="待办创建", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="todo_title", label="待办标题", type="text", required=True),
             dict(name="scope", label="负责范围", type="text", required=False),
             dict(name="due_time", label="截止时间", type="datetime", required=True),
             dict(name="priority", label="优先级", type="select", required=True,
                  options=opt([("HIGH", "高"), ("MID", "中"), ("LOW", "低")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=26, name="委托管理", source="原有", worker_type="OA_DELEGATE", module_code="delegate",
         form_code="delegate_rule", task_name="委托设置", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="delegatee", label="受托人", type="text", required=True),
             dict(name="start_date", label="委托开始", type="date", required=True),
             dict(name="end_date", label="委托结束", type="date", required=True),
             dict(name="matters", label="委托事项", type="textarea", required=True),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    # ---- 新增补充 17 ----
    dict(seq=27, name="门户首页", source="新增补充", worker_type="OA_PORTAL", module_code="portal",
         form_code="portal_layout", task_name="门户布局配置", biz_type="", submit_policy="LEDGER", derived=False,
         fields=[
             dict(name="workbench_name", label="工作台名称", type="text", required=True),
             dict(name="cards", label="卡片配置", type="textarea", required=True, placeholder="JSON/文本描述卡片组合"),
             dict(name="theme", label="主题", type="select", required=False,
                  options=opt([("DEFAULT", "默认"), ("DARK", "暗色"), ("LIGHT", "浅色")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=28, name="审批中心", source="新增补充", worker_type="OA_APPROVAL", module_code="approval",
         form_code="approval_launch", task_name="审批发起", biz_type="", submit_policy="LEDGER", derived=False,
         fields=[
             dict(name="matter", label="审批事项", type="text", required=True),
             dict(name="approval_type", label="审批类型", type="select", required=True,
                  options=opt([("GENERAL", "一般审批"), ("CONTRACT", "合同"), ("BUDGET", "预算"), ("HR", "人事")])),
             dict(name="urgency", label="紧急程度", type="select", required=True,
                  options=opt([("HIGH", "高"), ("MID", "中"), ("LOW", "低")])),
             dict(name="description", label="说明", type="textarea", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=29, name="人事管理", source="新增补充", worker_type="OA_HR", module_code="hr",
         form_code="hr_register", task_name="人事登记", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="事项类型", type="select", required=True,
                  options=opt([("ONBOARD", "入职登记"), ("TRANSFER", "异动转岗"), ("OFFBOARD", "离职交接")])),
             dict(name="emp_name", label="员工姓名", type="text", required=True),
             dict(name="dept", label="部门", type="text", required=True),
             dict(name="position", label="岗位", type="text", required=True),
             dict(name="effective_date", label="生效日期", type="date", required=True),
             dict(name="description", label="说明", type="textarea", required=False),
         ]),
    dict(seq=30, name="招聘管理", source="新增补充", worker_type="OA_RECRUIT", module_code="recruit",
         form_code="recruit_req", task_name="招聘需求", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="position", label="需求岗位", type="text", required=True),
             dict(name="headcount", label="编制数", type="number", required=True),
             dict(name="dept", label="需求部门", type="text", required=True),
             dict(name="expect_date", label="期望到岗", type="date", required=False),
             dict(name="duty", label="岗位职责", type="textarea", required=True),
         ]),
    dict(seq=31, name="培训管理", source="新增补充", worker_type="OA_TRAIN", module_code="train",
         form_code="train_signup", task_name="培训报名", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="train_name", label="培训名称", type="text", required=True),
             dict(name="course_type", label="课程类型", type="select", required=True,
                  options=opt([("ONLINE", "线上"), ("OFFLINE", "线下"), ("MIXED", "混合")])),
             dict(name="lecturer", label="讲师", type="text", required=False),
             dict(name="start_date", label="开始日期", type="date", required=True),
             dict(name="signup_note", label="报名说明", type="textarea", required=False),
         ]),
    dict(seq=32, name="知识管理", source="新增补充", worker_type="OA_KNOWLEDGE", module_code="knowledge",
         form_code="knowledge_upload", task_name="知识上传", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="knowledge_title", label="知识标题", type="text", required=True),
             dict(name="category", label="知识分类", type="select", required=True,
                  options=opt([("CASE", "经验案例"), ("POLICY", "制度"), ("SKILL", "技能"), ("FAQ", "问答")])),
             dict(name="tags", label="标签", type="text", required=False, placeholder="逗号分隔"),
             dict(name="summary", label="摘要", type="textarea", required=True),
             dict(name="file", label="文件", type="file", required=True),
         ]),
    dict(seq=33, name="资产管理", source="新增补充", worker_type="OA_ASSET", module_code="asset",
         form_code="asset_apply", task_name="资产办理", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="事项类型", type="select", required=True,
                  options=opt([("STOCK", "入库登记"), ("USE", "领用"), ("REPAIR", "维修"), ("SCRAP", "报废")])),
             dict(name="asset_name", label="资产名称", type="text", required=True),
             dict(name="quantity", label="数量", type="number", required=True),
             dict(name="dept", label="使用部门", type="text", required=False),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=34, name="合同管理", source="新增补充", worker_type="OA_CONTRACT", module_code="contract",
         form_code="contract_draft", task_name="合同起草", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="contract_name", label="合同名称", type="text", required=True),
             dict(name="counterparty", label="相对方", type="text", required=True),
             dict(name="amount", label="合同金额", type="amount", required=True),
             dict(name="start_date", label="开始日期", type="date", required=True),
             dict(name="end_date", label="结束日期", type="date", required=True),
             dict(name="description", label="说明", type="textarea", required=False),
         ]),
    dict(seq=35, name="任务督办", source="新增补充", worker_type="OA_TASK", module_code="task",
         form_code="task_supervise", task_name="督办立项", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="督办事项", type="text", required=True),
             dict(name="org_unit", label="责任单位", type="text", required=True),
             dict(name="owner", label="责任人", type="text", required=True),
             dict(name="due_date", label="截止日期", type="date", required=True),
             dict(name="level", label="紧急程度", type="select", required=True,
                  options=opt([("RED", "红"), ("YELLOW", "黄"), ("GREEN", "绿")])),
             dict(name="description", label="说明", type="textarea", required=False),
         ]),
    dict(seq=36, name="客户管理", source="新增补充", worker_type="OA_CRM", module_code="crm",
         form_code="crm_add", task_name="客户建档", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="customer_name", label="客户名称", type="text", required=True),
             dict(name="tier", label="客户分层", type="select", required=True,
                  options=opt([("STRATEGIC", "战略"), ("KEY", "重要"), ("NORMAL", "普通")])),
             dict(name="contact", label="联系人", type="text", required=False),
             dict(name="contact_phone", label="联系电话", type="text", required=False),
             dict(name="stage", label="商机阶段", type="select", required=False,
                  options=opt([("LEAD", "线索"), ("TALK", "洽谈"), ("WON", "成交"), ("LOST", "流失")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=37, name="财务管理", source="新增补充", worker_type="OA_FINANCE", module_code="finance",
         form_code="finance_apply", task_name="财务事项申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="matter", label="事项类型", type="select", required=True,
                  options=opt([("BUDGET", "预算编制"), ("PAY", "付款申请"), ("RECON", "对账")])),
             dict(name="amount", label="金额", type="amount", required=True),
             dict(name="description", label="说明", type="textarea", required=True),
             dict(name="project", label="关联项目", type="text", required=False),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=38, name="采购管理", source="新增补充", worker_type="OA_PURCHASE", module_code="purchase",
         form_code="purchase_req", task_name="采购申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="item", label="采购物品", type="text", required=True),
             dict(name="quantity", label="数量", type="number", required=True),
             dict(name="budget", label="预算金额", type="amount", required=True),
             dict(name="expect_arrive", label="期望到货", type="date", required=False),
             dict(name="purchase_note", label="采购说明", type="textarea", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
    dict(seq=39, name="问卷调查", source="新增补充", worker_type="OA_SURVEY", module_code="survey",
         form_code="survey_launch", task_name="问卷发起", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="survey_title", label="问卷标题", type="text", required=True),
             dict(name="purpose", label="调研目的", type="textarea", required=True),
             dict(name="scope", label="发放范围", type="textarea", required=True, placeholder="指定部门/人员"),
             dict(name="anonymous", label="匿名方式", type="select", required=True,
                  options=opt([("ANON", "匿名"), ("REALNAME", "实名")])),
             dict(name="due_date", label="截止日期", type="date", required=False),
         ]),
    dict(seq=40, name="安全与权限", source="新增补充", worker_type="OA_SECURITY", module_code="security",
         form_code="security_apply", task_name="权限申请", biz_type="", submit_policy="LEDGER", derived=True,
         fields=[
             dict(name="apply_type", label="申请类型", type="select", required=True,
                  options=opt([("ROLE", "角色授权"), ("DATA", "数据权限"), ("ACCOUNT", "账号开通")])),
             dict(name="target", label="申请对象", type="text", required=True),
             dict(name="scope", label="权限范围", type="textarea", required=True),
             dict(name="reason", label="理由", type="textarea", required=True),
         ]),
    dict(seq=41, name="系统集成", source="新增补充", worker_type="OA_INTEGRATION", module_code="integration",
         form_code="integration_req", task_name="集成申请", biz_type="", submit_policy="LEDGER", derived=False,
         fields=[
             dict(name="integration_type", label="集成类型", type="select", required=True,
                  options=opt([("SSO", "单点登录"), ("SYNC", "组织同步"), ("SIGN", "电子签章"), ("MSG", "消息推送")])),
             dict(name="system_name", label="系统名称", type="text", required=True),
             dict(name="description", label="说明", type="textarea", required=True),
         ]),
    dict(seq=42, name="数据分析", source="新增补充", worker_type="OA_ANALYTICS", module_code="analytics",
         form_code="analytics_board", task_name="看板配置", biz_type="", submit_policy="LEDGER", derived=False,
         fields=[
             dict(name="board_name", label="看板名称", type="text", required=True),
             dict(name="metrics", label="指标范围", type="textarea", required=True, placeholder="描述关注指标"),
             dict(name="push_role", label="推送角色", type="select", required=False,
                  options=opt([("LEADER", "领导"), ("MANAGER", "中层"), ("STAFF", "全员")])),
             dict(name="remark", label="备注", type="textarea", required=False),
         ]),
    dict(seq=43, name="移动办公", source="新增补充", worker_type="OA_MOBILE", module_code="mobile",
         form_code="mobile_checkin", task_name="移动签到", biz_type="", submit_policy="LEDGER", derived=False,
         fields=[
             dict(name="matter", label="事项类型", type="select", required=True,
                  options=opt([("CHECKIN", "移动签到"), ("FEEDBACK", "问题反馈"), ("SUGGEST", "功能建议")])),
             dict(name="location", label="位置", type="text", required=False, placeholder="定位地址"),
             dict(name="description", label="说明", type="textarea", required=True),
             dict(name="attachments", label="附件", type="file", required=False),
         ]),
]


# ---------------------------------------------------------------------------
# schema / receipt 构造
# ---------------------------------------------------------------------------
def build_schema(fields):
    props = {}
    required = []
    for f in fields:
        t = f["type"]
        if t not in TYPE_WHITELIST:
            raise ValueError("字段类型越白名单: %s -> %s" % (f["name"], t))
        p = {"type": BASE_JSON_TYPE[t], "title": f["label"], "x-type": t}
        if f.get("placeholder"):
            p["description"] = f["placeholder"]
        if t == "multiselect":
            p["items"] = {"type": "string"}
        if t in ("select", "multiselect") and f.get("options"):
            p["enum"] = [o["value"] for o in f["options"]]
            p["enumNames"] = [o["label"] for o in f["options"]]
        props[f["name"]] = p
        if f.get("required"):
            required.append(f["name"])
    return {
        "type": "object",
        "properties": props,
        "required": required,
        "additionalProperties": False,
    }


def build_receipt(name, fields):
    """回执结构必须与 `FormSubmissionService.renderReceipt` 认的契约一致。

    2026-10-17 实测缺陷：原先写的是 `display_fields: [name, ...]`，而后端只认
    `lines: [{label, value}]`（value 支持 `${field}` 占位）⇒ `lines` 恒为空 ⇒
    43 张表单的「声明式结果卡」全部退化成兜底回执「提交单号」，配置等于没生效。
    这里按 display 前 4 个字段生成 lines，label 取字段 title（与表单库同源）。
    """
    show = fields[:4]
    return {
        "title": "%s提交成功" % name,
        "lines": [{"label": f["label"], "value": "${%s}" % f["name"]} for f in show],
        "buttons": ["查看详情", "返回"],
    }


# ---------------------------------------------------------------------------
# SQL 转义
# ---------------------------------------------------------------------------
def sql_str(s):
    """MySQL 字符串字面量转义（单引号、反斜杠）。"""
    if s is None:
        return "NULL"
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_json(obj):
    """JSON 字面量，且**必须 Flyway 安全**：不能出现 `${`。

    2026-10-17 踩坑：回执里要用 `${field}` 做占位（`FormSubmissionService.substitute`），
    直接把 `$` 写进迁移文件，Flyway 会把它当**自己的占位符**解析并在启动时直接失败：
    `No value provided for placeholder: ${doc_kind}` —— 后端起不来，且报错完全没提是哪个文件。
    这里把 `$` 先写成 `§`，再用 `REPLACE(..., '§', CHAR(36))` 在库里还原，
    迁移文本里就不会出现 `${` 这个序列。
    """
    raw = json.dumps(obj, ensure_ascii=False).replace("$", "§")
    return "REPLACE(%s, '§', CHAR(36))" % sql_str(raw)


# ---------------------------------------------------------------------------
# 主生成
# ---------------------------------------------------------------------------
def check_xlsx():
    """读取 xlsx 校验 43 行，返回 (name_by_seq, source_by_seq)。"""
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    name_by_seq, src_by_seq = {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        seq = int(row[0])
        name_by_seq[seq] = str(row[1]).strip()
        src_by_seq[seq] = str(row[3]).strip()
    return name_by_seq, src_by_seq


def main():
    name_by_seq, src_by_seq = check_xlsx()
    assert len(name_by_seq) == 43, "xlsx 模块数应为 43，实际 %d" % len(name_by_seq)

    # 校验生成器内 MODULES 与 xlsx 一致（序号/名称/来源），但不依赖 xlsx 内容生成字段
    assert len(MODULES) == 43, "MODULES 数量应为 43，实际 %d" % len(MODULES)
    for m in MODULES:
        assert m["seq"] in name_by_seq, "MODULES seq=%d 在 xlsx 缺失" % m["seq"]
        assert m["name"] == name_by_seq[m["seq"]], \
            "seq=%d 名称不一致: 生成器=%s xlsx=%s" % (m["seq"], m["name"], name_by_seq[m["seq"]])
        assert m["source"] == src_by_seq[m["seq"]], \
            "seq=%d 来源不一致: 生成器=%s xlsx=%s" % (m["seq"], m["source"], src_by_seq[m["seq"]])
        # 字段数 3..6
        nf = len(m["fields"])
        assert 3 <= nf <= 6, "seq=%d 字段数 %d 不在 [3,6]" % (m["seq"], nf)
        # form_code 全局唯一
    codes = [m["form_code"] for m in MODULES]
    assert len(set(codes)) == 43, "form_code 不唯一"
    module_codes = [m["module_code"] for m in MODULES]
    assert len(set(module_codes)) == 43, "module_code 不唯一"

    lines = []
    lines.append("-- V86 · OA 43 个功能模块表单定义种子（docs/52 P5 首批铺开）")
    lines.append("--")
    lines.append("-- 由 scripts/_gen_oa_form_seed.py 生成，**勿手改**；复跑：")
    lines.append("--     python scripts/_gen_oa_form_seed.py")
    lines.append("-- 幂等：所有 INSERT 均带 WHERE NOT EXISTS / NOT EXISTS 子查询，可重复执行。")
    lines.append("-- 不改 V85（form_template 等四表结构以此文件为准），不 DELETE 任何业务数据。")
    lines.append("--")
    lines.append("-- 约定：biz_type 仅请假=LEAVE，其余空；submit_policy 仅请假=APPROVAL+RESUME，其余=LEDGER。")
    lines.append("-- schema_json 用 x-type 承载白名单控件类型（text/textarea/number/amount/date/datetime/")
    lines.append("-- select/multiselect/file/boolean），select/multiselect 用 enum+enumNames 内联静态选项。")
    lines.append("")

    # ---- 1) 43 条 form_template（平台默认 tenant_id=0）----
    lines.append("-- ── 1) form_template：43 个模块平台默认定义（tenant_id=0, created_by=0）──")
    for m in MODULES:
        schema = build_schema(m["fields"])
        receipt = build_receipt(m["name"], m["fields"])
        biz = sql_str(m["biz_type"]) if m["biz_type"] else "NULL"
        lines.append(
            "INSERT INTO `form_template` "
            "(`tenant_id`, `code`, `name`, `module_code`, `biz_type`, `schema_json`, "
            "`layout_json`, `submit_policy`, `receipt_json`, `version`, `status`, `created_by`, `created_at`, `updated_at`)"
        )
        lines.append("SELECT 0, %s, %s, %s, %s, %s, NULL, %s, %s, 1, 'ACTIVE', 0, NOW(6), NOW(6)"
                     % (sql_str(m["form_code"]), sql_str(m["name"]), sql_str(m["module_code"]),
                        biz, sql_json(schema), sql_str(m["submit_policy"]), sql_json(receipt)))
        lines.append("WHERE NOT EXISTS (SELECT 1 FROM `form_template` x "
                     "WHERE x.`tenant_id` = 0 AND x.`code` = %s);" % sql_str(m["form_code"]))
        lines.append("")

    # ---- 2) 按租户回填 form_template（仿 V83 CROSS JOIN + WHERE NOT EXISTS）----
    code_list = ", ".join(sql_str(c) for c in codes)
    lines.append("-- ── 2) 按租户回填 form_template：以平台定义为唯一数据源，逐租户复制 ──")
    lines.append("INSERT INTO `form_template`")
    lines.append("    (`tenant_id`, `code`, `name`, `module_code`, `biz_type`, `schema_json`,")
    lines.append("     `layout_json`, `submit_policy`, `receipt_json`, `version`, `status`, `created_by`, `created_at`, `updated_at`)")
    lines.append("SELECT t.`id`, m.`code`, m.`name`, m.`module_code`, m.`biz_type`, m.`schema_json`,")
    lines.append("       m.`layout_json`, m.`submit_policy`, m.`receipt_json`, 1, 'ACTIVE', 0, NOW(6), NOW(6)")
    lines.append("FROM `sys_tenant` t")
    lines.append("CROSS JOIN (")
    lines.append("    SELECT `code`, `name`, `module_code`, `biz_type`, `schema_json`, `layout_json`, `submit_policy`, `receipt_json`")
    lines.append("    FROM `form_template` WHERE `tenant_id` = 0 AND `code` IN (%s)" % code_list)
    lines.append(") m")
    lines.append("WHERE t.`id` > 0")
    lines.append("  AND NOT EXISTS (SELECT 1 FROM `form_template` x WHERE x.`tenant_id` = t.`id` AND x.`code` = m.`code`);")
    lines.append("")

    # ---- 3) form_template_version：每个 form_template 行一条 version=1 快照 ----
    lines.append("-- ── 3) form_template_version：每个定义（平台+租户回填）各一条 version=1 快照 ──")
    lines.append("INSERT INTO `form_template_version` (`template_id`, `version`, `schema_json`, `layout_json`, `receipt_json`, `created_by`)")
    lines.append("SELECT f.`id`, 1, f.`schema_json`, f.`layout_json`, f.`receipt_json`, 0")
    lines.append("FROM `form_template` f")
    lines.append("WHERE f.`code` IN (%s)" % code_list)
    lines.append("  AND NOT EXISTS (SELECT 1 FROM `form_template_version` v WHERE v.`template_id` = f.`id` AND v.`version` = 1);")
    lines.append("")

    # ---- 4) worker_capability：每个模块 1 条能力声明（模板 + 各租户实例）----
    lines.append("-- ── 4) worker_capability：43 模块各 1 条能力声明，从 agent_worker 派生（模板+实例）──")
    lines.append("--    JOIN 条件 t.tenant_id = w.tenant_id AND t.module_code = <稳定码>，保证同租户匹配；")
    lines.append("--    worker_type→module_code 映射固化于下方 CASE/WHERE；全部 WHERE NOT EXISTS 幂等。")
    for m in MODULES:
        approval = sql_str(m["biz_type"]) if m["biz_type"] else "NULL"
        lines.append("-- %d %s (%s → %s)" % (m["seq"], m["name"], m["worker_type"], m["module_code"]))
        lines.append("INSERT INTO `worker_capability`")
        lines.append("    (`tenant_id`, `worker_id`, `task_code`, `task_name`, `form_code`, `tool_names_json`, `approval_biz_type`, `sort`, `enabled`)")
        lines.append("SELECT w.`tenant_id`, w.`id`, %s, %s, t.`code`, NULL, %s, 0, 1"
                     % (sql_str(m["form_code"]), sql_str(m["task_name"]), approval))
        lines.append("FROM `agent_worker` w")
        lines.append("JOIN `form_template` t ON t.`tenant_id` = w.`tenant_id` AND t.`module_code` = %s"
                     % sql_str(m["module_code"]))
        lines.append("WHERE w.`worker_type` = %s" % sql_str(m["worker_type"]))
        lines.append("  AND NOT EXISTS (SELECT 1 FROM `worker_capability` c"
                     " WHERE c.`tenant_id` = w.`tenant_id` AND c.`worker_id` = w.`id` AND c.`task_code` = %s);"
                     % sql_str(m["form_code"]))
        lines.append("")

    out = "\n".join(lines) + "\n"
    os.makedirs(os.path.dirname(OUT_SQL), exist_ok=True)
    with open(OUT_SQL, "w", encoding="utf-8") as fh:
        fh.write(out)
    print("OK 生成 %s" % OUT_SQL)
    print("  模块数=%d  字段数区间=%s..%s  语句数(含空行)=%d"
          % (len(MODULES),
             min(len(m["fields"]) for m in MODULES),
             max(len(m["fields"]) for m in MODULES),
             len(lines)))
    # 统计 derived
    print("  严格推导模块=%d  合理默认模块=%d"
          % (sum(1 for m in MODULES if m["derived"]), sum(1 for m in MODULES if not m["derived"])))


if __name__ == "__main__":
    main()
