# -*- coding: utf-8 -*-
"""为公式单元格补写「缓存值」，使非 Excel 预览器也能看到统计结果。

背景：openpyxl 写出的公式只带公式、不带缓存值，Excel/WPS 打开时才计算，
而浏览器预览 / pandas / data_only=True 读到的都是空。本机既无 LibreOffice，
包源也没有 formulas 引擎，故按「先算后注入 XML」补齐缓存值。

注入值不是手写常量：由同一份工作簿的用例行现算得出，注入后再读回断言，闭环可验。
"""
import re
import shutil
import zipfile

from openpyxl import load_workbook

XLSX = "C:/Users/刘尖尖/WorkBuddy/aioa/docs/41-项目管理与项目仓库_手工功能测试用例.xlsx"
TMP = XLSX + ".tmp"

# 【执行结果】列的取值全集
DONE = ("通过", "失败", "阻塞")


def last_case_row(ws):
    """用例数据区末行：以「用例ID」列能匹配 ^(PM|GR)-\\d+$ 的最大行号为锚点。

    不能用「A 列非空计数」——统计区的标签也在 A 列，会把锚点算高。
    """
    rows = [r for r in range(3, ws.max_row + 1)
            if isinstance(ws.cell(row=r, column=1).value, str)
            and re.match(r"^(PM|GR)-\d+$", ws.cell(row=r, column=1).value)]
    assert rows, "未找到任何用例行"
    return max(rows)


def metrics(ws, last_row):
    """按真实用例行现算 7 个指标，与表内公式同口径。"""
    total = sum(1 for r in range(3, last_row + 1) if ws.cell(row=r, column=1).value)
    pos = sum(1 for r in range(3, last_row + 1) if ws.cell(row=r, column=3).value == "正向")
    neg = sum(1 for r in range(3, last_row + 1) if ws.cell(row=r, column=3).value == "反向")
    executed = sum(1 for r in range(3, last_row + 1) if ws.cell(row=r, column=9).value in DONE)
    passed = sum(1 for r in range(3, last_row + 1) if ws.cell(row=r, column=9).value == "通过")
    pending = sum(1 for r in range(3, last_row + 1) if ws.cell(row=r, column=9).value == "未执行")
    rate = "" if executed in ("", 0) else passed / executed
    return total, pos, neg, executed, passed, rate, pending


wb = load_workbook(XLSX)
targets = {}  # sheet_xml_name -> {cell_ref: (value, is_string)}

for xml_name, ws in (("xl/worksheets/sheet2.xml", wb["项目管理用例"]),
                     ("xl/worksheets/sheet3.xml", wb["项目与仓库用例"])):
    last = last_case_row(ws)
    anchor = last + 2
    vals = metrics(ws, last)
    cells = {}
    for i, v in enumerate(vals):
        ref = f"B{anchor + 1 + i}"
        cells[ref] = (v, isinstance(v, str))
    targets[xml_name] = cells
    print(f"{ws.title}: 数据行 3..{last}  统计区起点行 {anchor + 1}  值={vals}")

# ---- 注入：只替换目标单元格里的空 <v></v>，其余 XML 原样保留 ----
zin = zipfile.ZipFile(XLSX, "r")
zout = zipfile.ZipFile(TMP, "w", zipfile.ZIP_DEFLATED)
patched = 0
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in targets:
        xml = data.decode("utf-8")
        for ref, (val, is_str) in targets[item.filename].items():
            if is_str:
                # 文本结果：标 t="str"，空字符串写空 <v/>
                pat = re.compile(r'(<c r="%s"([^>]*?)>)' % ref)
                xml, n = pat.subn(lambda m: f'<c r="{ref}"{m.group(2)} t="str">', xml, count=1)
                assert n == 1, f"未定位到 {ref} 的开标签"
            cell_pat = re.compile(r'(<c r="%s"[^>]*>)(<f>.*?</f>)<v></v>(</c>)' % ref, re.S)
            xml, n = cell_pat.subn(lambda m: m.group(1) + m.group(2) + f"<v>{val}</v>" + m.group(3), xml, count=1)
            assert n == 1, f"未注入 {ref}"
            patched += 1
        data = xml.encode("utf-8")
    zout.writestr(item, data)
zin.close()
zout.close()
shutil.move(TMP, XLSX)
print("patched cells:", patched)

# ---- 闭环校验：以 data_only=True 读回，断言缓存值与现算一致 ----
wb2 = load_workbook(XLSX, data_only=True)
ok = True
for name in ("项目管理用例", "项目与仓库用例"):
    ws = wb2[name]
    ws0 = wb[name]
    last = last_case_row(ws0)
    anchor = last + 2
    expect = metrics(ws0, last)
    got = tuple(ws.cell(row=anchor + 1 + i, column=2).value for i in range(7))
    norm_got = tuple("" if g is None else g for g in got)
    print(f"[{name}] expect={expect}")
    print(f"[{name}] cached={norm_got}")
    if norm_got != tuple(expect):
        ok = False
        print(f"[{name}] !! 缓存值与现算不一致")
print("VERIFY:", "PASS" if ok else "FAIL")
raise SystemExit(0 if ok else 1)
